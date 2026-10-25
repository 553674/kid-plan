import streamlit as st
from docx import Document
import openai, io, re

DEFAULT_TEMPLATE = """活动名称：{{活动名称}}
活动目标：
1. 认知目标：{{认知}}
2. 情感目标：{{情感}}
3. 技能目标：{{技能}}
活动重难点：
重点：{{重点}}
难点：{{难点}}
活动准备：
1. 经验准备：{{经验}}
2. 物质准备：{{物质}}
活动过程：
一、导入活动
1. {{导入1}}
2. {{导入2}}
二、基本活动
1. {{基本1}}
2. {{基本2}}
三、结束活动
1. {{结束1}}
2. {{结束2}}
活动延伸：
1. 区域延伸：{{区域}}
2. 生活延伸：{{生活}}
活动建议与反思：{{反思}}"""

DAYS = ["星期一","星期二","星期三","星期四","星期五"]
SLOTS = ["晨间活动","早餐","教学活动","户外自主游戏","户外运动",
         "盥洗","如厕","喝水","餐前活动","午餐","午睡","区域活动","安全活动"]

def gen(day, slot, name, template, api_key):
    client = openai.OpenAI(api_key=api_key, base_url="https://api.moonshot.cn/v1")
    prompt = f"""严格按下面格式模板输出，替换所有{{{{}}}}占位符为实际内容。
不要输出模板外任何文字，不要用Markdown符号（# * -）。

模板：
{template}

要生成：
星期：{day}
时段：{slot}
活动名称：{name}
"""
    r = client.chat.completions.create(
        model="moonshot-v1-8k",
        messages=[{"role":"user","content":prompt}],
        temperature=0.4)
    return r.choices[0].message.content.strip()

def parse_week(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    result = {}
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            row_text = " ".join(c.text for c in cells)
            hit_slot = next((s for s in SLOTS if s in row_text), None)
            if not hit_slot: continue
            for ci, cell in enumerate(cells):
                t = cell.text.strip()
                if t and len(t) < 60 and ("：" in t or ":" in t):
                    day = DAYS[ci] if ci < 5 else f"列{ci}"
                    result[(day, hit_slot)] = t
    return result

def parse_day(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    pos = []
    seen = set()
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                tc_id = id(cell._tc)
                if tc_id in seen: continue
                seen.add(tc_id)
                t = cell.text
                for slot in SLOTS:
                    if slot in t:
                        pos.append({"key":f"{ti}-{ri}-{ci}-{slot}",
                                    "table":ti,"row":ri,"col":ci,
                                    "slot":slot,"text":t})
                        break
    return doc, pos

def write_back(doc, pos, changes, template, api_key):
    for p in pos:
        if p["key"] not in changes: continue
        day, slot, name = changes[p["key"]]
        new_text = gen(day, slot, name, template, api_key)
        cell = doc.tables[p["table"]].rows[p["row"]].cells[p["col"]]
        old = cell.text
        pattern = rf"{re.escape(slot)}.*?(?=(?:{'|'.join(map(re.escape,SLOTS))})|$)"
        cell.text = re.sub(pattern, new_text, old, flags=re.S)
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf

st.set_page_config(page_title="周计划→一日计划", layout="wide")
st.title("📅 一键生成一日计划内容")

api_key = st.text_input("Kimi API Key", type="password")

with st.expander("⚙️ 自定义输出格式模板（改这里AI就按新格式生成）"):
    template = st.text_area("模板", DEFAULT_TEMPLATE, height=380)

week_file = st.file_uploader("① 上传 周计划.docx", type="docx")
day_file  = st.file_uploader("② 上传 一日计划.docx", type="docx")

if week_file and day_file:
    week_map = parse_week(week_file.read())
    day_doc, day_pos = parse_day(day_file.read())

    st.subheader("勾选要修改的时段")
    changes = {}
    for (day, slot), name in sorted(week_map.items()):
        c1, c2, c3 = st.columns([0.3, 2, 3])
        with c1: chk = st.checkbox("", key=f"{day}-{slot}")
        with c2: st.text(f"{day}/{slot}")
        with c3: new_name = st.text_input("活动名", name,
                    key=f"n-{day}-{slot}", label_visibility="collapsed")
        if chk:
            for p in day_pos:
                if p["slot"] == slot:
                    changes[p["key"]] = (day, slot, new_name)
                    break

    if st.button("🚀 生成并下载", type="primary"):
        if not api_key: st.error("请填 Kimi API Key")
        elif not changes: st.warning("没勾选")
        else:
            with st.spinner("生成中..."):
                buf = write_back(day_doc, day_pos, changes, template, api_key)
            st.success("完成，未勾选内容未改动")
            st.download_button("⬇️ 下载新一日计划", buf,
                file_name="新_"+day_file.name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
